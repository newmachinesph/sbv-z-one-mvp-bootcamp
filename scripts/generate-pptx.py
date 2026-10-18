"""
Z1 Forge — Light-mode presentation with blueprint imagery + Z1 logo.

Images: public/blueprint/1-13.jpg (896x1200, 3:4 portrait)
Logo:   public/z1-logo-orange.png (500x585, amber Z1 mark)

Fonts: Manrope (headings), DM Sans (body)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG  = lambda n: os.path.join(BASE, "public", "blueprint", f"{n}.jpg")
LOGO = os.path.join(BASE, "public", "z1-logo-orange.png")

# ── Colour tokens ────────────────────────────────────────────
AMBER       = RGBColor(0xF5, 0x9E, 0x0B)
DARK_AMBER  = RGBColor(0xD9, 0x77, 0x06)
AMBER_50    = RGBColor(0xFF, 0xFB, 0xEB)
AMBER_100   = RGBColor(0xFE, 0xF3, 0xC7)
AMBER_200   = RGBColor(0xFD, 0xE6, 0x8A)
TEXT        = RGBColor(0x11, 0x18, 0x27)
SUBTEXT     = RGBColor(0x4B, 0x55, 0x63)
MUTED       = RGBColor(0x9C, 0xA3, 0xAF)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BLACK       = RGBColor(0x00, 0x00, 0x00)
CARD_BG     = RGBColor(0xF9, 0xFA, 0xFB)
CARD_BORDER = RGBColor(0xE5, 0xE7, 0xEB)
DARK_BG     = RGBColor(0x0A, 0x0A, 0x0A)
LIGHT_TEXT  = RGBColor(0xD1, 0xD5, 0xDB)

H = "Manrope"   # headings
B = "DM Sans"   # body

W = Inches(13.333)
HH = Inches(7.5)
TOTAL = 13

# Split layout: content on left, image panel on right
CONTENT_W = Inches(7.8)
IMG_LEFT  = Inches(8.0)
IMG_W     = Inches(5.333)

prs = Presentation()
prs.slide_width = W; prs.slide_height = HH


# ═══════════════════════════════════════════════════════════════
#  Core helpers
# ═══════════════════════════════════════════════════════════════

def _alpha(shape, opacity):
    """Set fill opacity via XML. 1.0 = fully opaque, 0.0 = transparent."""
    for sf in shape._element.findall('.//' + qn('a:solidFill')):
        sr = sf.find(qn('a:srgbClr'))
        if sr is not None:
            for old in sr.findall(qn('a:alpha')): sr.remove(old)
            a = etree.SubElement(sr, qn('a:alpha'))
            a.set('val', str(int(opacity * 100000)))


def set_bg(slide, color=WHITE):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def bar_top(slide):
    """Amber accent bar across top."""
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.05))
    s.fill.solid(); s.fill.fore_color.rgb = AMBER; s.line.fill.background()


def add_logo(slide, left=Inches(0.7), top=Inches(0.25), height=Inches(0.4)):
    """Place Z1 logo mark. Positioned below the accent bar with padding."""
    if os.path.exists(LOGO):
        slide.shapes.add_picture(LOGO, left, top, height=height)


def full_bg(slide, img_num):
    """Full-bleed background image, center-cropped for 16:9."""
    path = IMG(img_num)
    if not os.path.exists(path): return
    # Images are 896x1200 (3:4 portrait). Slide is 16:9 landscape.
    # Scale to fill width, then center-crop vertically.
    pic = slide.shapes.add_picture(path, 0, 0, width=W)
    # Calculate actual height at this width
    img_ratio = 1200 / 896  # ~1.339
    display_h = int(W * img_ratio)
    pic.height = display_h
    pic.top = -(display_h - HH) // 2  # center vertically
    # Send to back of z-order
    el = pic._element; el.getparent().remove(el)
    slide.shapes._spTree.insert(2, el)


def overlay(slide, opacity=0.55):
    """Dark semi-transparent overlay."""
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, HH)
    s.fill.solid(); s.fill.fore_color.rgb = DARK_BG
    _alpha(s, opacity); s.line.fill.background()


def right_image(slide, img_num):
    """Portrait image filling the right panel with soft left-edge fade."""
    path = IMG(img_num)
    if not os.path.exists(path): return
    # Scale image to fill panel height (7.5"). Width = 7.5/1.339 = 5.6"
    pic = slide.shapes.add_picture(path, 0, 0, height=HH)
    # Right-align the image
    pic.left = W - pic.width
    # Soft white gradient fade on left edge
    fade_w = Inches(1.5)
    fade = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   pic.left - Inches(0.1), 0, fade_w, HH)
    fade.fill.gradient()
    fade.fill.gradient_stops[0].color.rgb = WHITE
    fade.fill.gradient_stops[0].position = 0.0
    fade.fill.gradient_stops[1].color.rgb = WHITE
    fade.fill.gradient_stops[1].position = 1.0
    # First stop opaque, second stop transparent
    for i, gs in enumerate(fade.fill.gradient_stops):
        sr = gs._element.find(qn('a:srgbClr'))
        if sr is not None:
            a = etree.SubElement(sr, qn('a:alpha'))
            a.set('val', '100000' if i == 0 else '0')
    fade.line.fill.background(); fade.rotation = 0.0


def page_num(slide, n, color=MUTED):
    tx(slide, Inches(12.0), Inches(7.05), Inches(1.0), Inches(0.3),
       f"{n:02d}", 9, color, font=B, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, fill_c=CARD_BG, border_c=CARD_BORDER):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = fill_c
    r.line.color.rgb = border_c; r.line.width = Pt(0.75)
    try: r.adjustments[0] = 0.02
    except: pass
    return r


def acard(slide, x, y, w, h):
    return card(slide, x, y, w, h, AMBER_50, AMBER_200)


def rule(slide, left, top, width, color=AMBER_200):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.02))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()


def tx(slide, left, top, width, height, text, size=16,
       color=TEXT, bold=False, font=B, align=PP_ALIGN.LEFT, sp=1.35):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    r = p.runs[0]; r.font.size = Pt(size); r.font.color.rgb = color
    r.font.name = font; r.font.bold = bold
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    if sp != 1.0: p.line_spacing = Pt(size * sp)
    return tb, tf


def label(slide, text, top=Inches(0.65), color=DARK_AMBER):
    tx(slide, Inches(1.5), top, Inches(5), Inches(0.25),
       text, 10, color, bold=True, font=H)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 1 — Cover  (full-bleed: 1.jpg forge anvil)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, BLACK)
full_bg(s, 1)
overlay(s, 0.50)
bar_top(s)
add_logo(s, Inches(0.9), Inches(0.5), Inches(0.55))

tx(s, Inches(0.9), Inches(2.2), Inches(8), Inches(1.0),
   "Z1 Forge", 58, WHITE, bold=True, font=H, sp=1.05)
tx(s, Inches(0.9), Inches(3.5), Inches(6), Inches(0.5),
   "by SBV", 26, AMBER, font=H)
tx(s, Inches(0.9), Inches(4.5), Inches(7), Inches(0.6),
   "From Idea to MVP in 12 Weeks", 24, WHITE, bold=True, font=H)
tx(s, Inches(0.9), Inches(5.4), Inches(7), Inches(0.8),
   "The collaborative forge where small teams build, ship, and launch real products with real users \u2014 in just 12 weeks.",
   14, LIGHT_TEXT, font=B, sp=1.7)

# Bottom amber strip
bb = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.05), W, Inches(0.45))
bb.fill.solid(); bb.fill.fore_color.rgb = AMBER; bb.line.fill.background()
tx(s, Inches(0.9), Inches(7.1), Inches(6), Inches(0.35),
   "z1.sbv.one  \u2022  Build. Ship. Validate. Iterate.", 12, BLACK, bold=True, font=H)
page_num(s, 1, RGBColor(0x66, 0x66, 0x66))


# ═══════════════════════════════════════════════════════════════
#  SLIDE 2 — Company Overview  (split: 2.jpg coworking)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 2)
add_logo(s); label(s, "COMPANY OVERVIEW")

tx(s, Inches(1.0), Inches(1.2), Inches(6.5), Inches(0.7),
   "Z1 Forge by SBV", 34, TEXT, bold=True, font=H)

tx(s, Inches(1.0), Inches(2.0), Inches(6.3), Inches(0.9),
   "A 12-week execution-driven build program designed to help founders, students, and professionals move from zero idea to a launched MVP with real users and traction.",
   13, SUBTEXT, font=B, sp=1.7)

# Statements
for i, (st, hl) in enumerate([
    ("We are not a classroom.", False),
    ("We are not a pitch competition.", False),
    ("We are a build-and-ship environment.", True),
]):
    y = Inches(3.2) + Inches(0.38) * i
    tx(s, Inches(1.0), y, Inches(6.3), Inches(0.3), st, 13,
       DARK_AMBER if hl else TEXT, bold=hl, font=H)

# Fact cards
facts = [
    ("PROGRAM TYPE", "MVP Accelerator /\nVenture Build"),
    ("FOCUS", "Zero \u2192 One\nExecution"),
    ("FORMAT", "Hybrid /\nCohort-Based"),
    ("DURATION", "12\nWeeks"),
]
cw = Inches(1.65); ch = Inches(1.8); gap = Inches(0.1); sy = Inches(4.6)
for i, (lbl, val) in enumerate(facts):
    x = Inches(1.0) + (cw + gap) * i
    acard(s, x, sy, cw, ch)
    tx(s, x + Inches(0.15), sy + Inches(0.15), cw - Inches(0.3), Inches(0.2),
       lbl, 8, DARK_AMBER, bold=True, font=H)
    tx(s, x + Inches(0.15), sy + Inches(0.5), cw - Inches(0.3), Inches(1.1),
       val, 15, TEXT, bold=True, font=H, sp=1.25)

page_num(s, 2)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 3 — About  (split: 3.jpg Makati skyline)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 3)
add_logo(s); label(s, "ABOUT Z1 FORGE")

tx(s, Inches(1.0), Inches(1.3), Inches(6.5), Inches(0.7),
   "Most founders don\u2019t fail\nbecause their ideas are bad.", 28, TEXT, bold=True, font=H, sp=1.25)

tx(s, Inches(1.0), Inches(2.6), Inches(6.5), Inches(0.5),
   "They fail because they never launch.", 28, DARK_AMBER, bold=True, font=H)

rule(s, Inches(1.0), Inches(3.4), Inches(2.5))

tx(s, Inches(1.0), Inches(3.7), Inches(6.3), Inches(1.1),
   "Z1 Forge was created to solve the execution gap \u2014 turning ideas into working products through structured weekly deliverables, mentorship, accountability, and real-world validation.",
   13, SUBTEXT, font=B, sp=1.7)

# Goal card
acard(s, Inches(1.0), Inches(5.2), Inches(6.5), Inches(1.2))
tx(s, Inches(1.2), Inches(5.3), Inches(2), Inches(0.2),
   "OUR GOAL", 9, DARK_AMBER, bold=True, font=H)
tx(s, Inches(1.2), Inches(5.6), Inches(6.0), Inches(0.6),
   "Help founders ship fast, validate early, and build products users actually want.",
   16, TEXT, bold=True, font=H, sp=1.3)

page_num(s, 3)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 4 — Vision  (full-bleed: 5.jpg rice terraces road)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, BLACK)
full_bg(s, 5)
overlay(s, 0.55)
bar_top(s)
add_logo(s, Inches(0.9), Inches(0.5), Inches(0.45))
label(s, "VISION", top=Inches(0.6), color=AMBER)

# Left accent bar
lb = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.8), Inches(0.06), Inches(3.5))
lb.fill.solid(); lb.fill.fore_color.rgb = AMBER; lb.line.fill.background()

tx(s, Inches(1.3), Inches(1.8), Inches(10), Inches(3.5),
   "To become the leading\nZero \u2192 One venture build ecosystem\nthat empowers execution-driven founders\nto launch real products with\nmeasurable traction.",
   32, WHITE, bold=True, font=H, sp=1.4)

rule(s, Inches(1.3), Inches(5.8), Inches(2.5), AMBER)
page_num(s, 4, RGBColor(0x88, 0x88, 0x88))


# ═══════════════════════════════════════════════════════════════
#  SLIDE 5 — Mission  (split: 4.jpg laptop)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 4)
add_logo(s); label(s, "MISSION")

tx(s, Inches(1.0), Inches(1.2), Inches(6.5), Inches(0.8),
   "To guide founders from\nidea to MVP through:", 24, TEXT, bold=True, font=H, sp=1.25)

mission = [
    "Structured execution frameworks",
    "Weekly build milestones",
    "Real-user validation",
    "Operator-led mentorship",
    "Global Demo Day exposure",
]
cw = Inches(6.5); ch = Inches(0.6); sy = Inches(2.5)
for i, item in enumerate(mission):
    y = sy + (ch + Inches(0.1)) * i
    card(s, Inches(1.0), y, cw, ch)
    # Amber number
    tx(s, Inches(1.2), y + Inches(0.1), Inches(0.4), Inches(0.4),
       f"0{i+1}", 12, DARK_AMBER, bold=True, font=H)
    tx(s, Inches(1.7), y + Inches(0.1), cw - Inches(0.9), Inches(0.4),
       item, 14, TEXT, font=B)

page_num(s, 5)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 6 — Core Values  (split: 6.jpg glass marbles)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 6)
add_logo(s); label(s, "CORE VALUES")

values = [
    ("Execution Over Perfection", "We ship fast and iterate based on feedback."),
    ("Real-World Validation", "Ideas are tested with users, not assumptions."),
    ("Accountability & Discipline", "Weekly deliverables drive progress."),
    ("Operator-Led Guidance", "Mentors are builders, not just speakers."),
    ("Zero \u2192 One Focus", "We prioritize launch and traction over theory."),
]
cw = Inches(6.5); ch = Inches(1.05); sy = Inches(1.2)
for i, (title, desc) in enumerate(values):
    y = sy + (ch + Inches(0.1)) * i
    card(s, Inches(1.0), y, cw, ch)

    # Number circle
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), y + Inches(0.22), Inches(0.5), Inches(0.5))
    circ.fill.solid(); circ.fill.fore_color.rgb = AMBER; circ.line.fill.background()
    tx(s, Inches(1.2), y + Inches(0.25), Inches(0.5), Inches(0.45),
       str(i+1), 16, WHITE, bold=True, font=H, align=PP_ALIGN.CENTER)

    tx(s, Inches(1.9), y + Inches(0.12), Inches(5.2), Inches(0.3),
       title, 14, TEXT, bold=True, font=H)
    tx(s, Inches(1.9), y + Inches(0.5), Inches(5.2), Inches(0.4),
       desc, 11, SUBTEXT, font=B)

page_num(s, 6)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 7 — What We Offer  (split: 7.jpg spiral staircase)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 7)
add_logo(s); label(s, "WHAT WE OFFER")

offerings = [
    ("12-Week Structured MVP Build Curriculum",
     "Step-by-step roadmap: validation \u2192 design \u2192 prototype \u2192 MVP \u2192 traction."),
    ("Weekly Build Sprints",
     "Clear deliverables every week to ensure measurable progress."),
    ("Expert Mentorship",
     "Selected operators provide tactical guidance in product, growth, tech, and execution."),
    ("Demo Day Exposure",
     "Founders showcase working MVPs to partners, investors, and ecosystem players."),
    ("Founder Visibility",
     "Selected teams receive spotlight campaigns and brand exposure."),
]
cw = Inches(6.5); ch = Inches(1.05); sy = Inches(1.2)
for i, (title, desc) in enumerate(offerings):
    y = sy + (ch + Inches(0.1)) * i
    acard(s, Inches(1.0), y, cw, ch)

    # Amber dot
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), y + Inches(0.18), Inches(0.12), Inches(0.12))
    dot.fill.solid(); dot.fill.fore_color.rgb = AMBER; dot.line.fill.background()

    tx(s, Inches(1.5), y + Inches(0.1), Inches(5.7), Inches(0.3),
       title, 13, TEXT, bold=True, font=H)
    tx(s, Inches(1.2), y + Inches(0.5), Inches(6.0), Inches(0.4),
       desc, 11, SUBTEXT, font=B, sp=1.4)

page_num(s, 7)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 8 — Target Participants  (split: 8.jpg glass spheres)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 8)
add_logo(s); label(s, "TARGET PARTICIPANTS")

tx(s, Inches(1.0), Inches(1.2), Inches(6.5), Inches(0.5),
   "Z1 Forge is built for:", 24, TEXT, bold=True, font=H)

participants = [
    ("Early-stage founders", "Ready to execute and build."),
    ("Students with startup ideas", "Turn projects into real products."),
    ("Professionals transitioning\ninto entrepreneurship", "Leverage expertise to build."),
    ("Builders who want accountability\nand structure", "Framework-driven idea \u2192 launch."),
]
cw = Inches(3.1); ch = Inches(1.15); gx = Inches(0.2); gy = Inches(0.15); sy = Inches(1.9)
for i, (title, desc) in enumerate(participants):
    col = i % 2; row = i // 2
    x = Inches(1.0) + (cw + gx) * col
    y = sy + (ch + gy) * row
    card(s, x, y, cw, ch)
    tx(s, x + Inches(0.2), y + Inches(0.1), cw - Inches(0.4), Inches(0.5),
       title, 12, TEXT, bold=True, font=H, sp=1.2)
    tx(s, x + Inches(0.2), y + Inches(0.7), cw - Inches(0.4), Inches(0.35),
       desc, 10, SUBTEXT, font=B)

# Industries
acard(s, Inches(1.0), Inches(4.7), Inches(6.5), Inches(1.0))
tx(s, Inches(1.2), Inches(4.8), Inches(2), Inches(0.2),
   "INDUSTRIES", 9, DARK_AMBER, bold=True, font=H)
tx(s, Inches(1.2), Inches(5.1), Inches(6.0), Inches(0.5),
   "Tech  \u2022  SaaS  \u2022  Marketplace  \u2022  EdTech\nFinTech  \u2022  AI Tools  \u2022  Scalable Digital Products",
   12, TEXT, bold=True, font=H, sp=1.4)

page_num(s, 8)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 9 — UVP  (split: 9.jpg hands reaching)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 9)
add_logo(s); label(s, "UNIQUE VALUE PROPOSITION")

tx(s, Inches(1.0), Inches(1.1), Inches(6.5), Inches(0.8),
   "Unlike traditional incubators that focus on theory and pitch decks, Z1 Forge is:",
   13, SUBTEXT, font=B, sp=1.7)

pillars = ["Execution-\ndriven", "Deliverable-\nbased", "Traction-\nfocused", "Built around\nMVP launch"]
pw = Inches(1.55); ph = Inches(1.3); pgap = Inches(0.1)
for i, pillar in enumerate(pillars):
    x = Inches(1.0) + (pw + pgap) * i
    y = Inches(2.2)
    # Amber top accent
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, pw, Inches(0.04))
    bar.fill.solid(); bar.fill.fore_color.rgb = AMBER; bar.line.fill.background()
    card(s, x, y + Inches(0.04), pw, ph)
    tx(s, x, y + Inches(0.25), pw, Inches(0.9),
       pillar, 13, TEXT, bold=True, font=H, align=PP_ALIGN.CENTER, sp=1.25)

# Metrics
tx(s, Inches(1.0), Inches(3.9), Inches(5), Inches(0.3),
   "We measure success by:", 13, TEXT, bold=True, font=H)

metrics = ["\u2714  Working product", "\u2714  Real users", "\u2714  Early traction", "\u2714  Founder momentum"]
for i, m in enumerate(metrics):
    col = i % 2; row = i // 2
    x = Inches(1.0) + Inches(3.3) * col
    y = Inches(4.4) + Inches(0.5) * row
    acard(s, x, y, Inches(3.0), Inches(0.4))
    tx(s, x, y + Inches(0.04), Inches(3.0), Inches(0.3),
       m, 11, TEXT, font=B, align=PP_ALIGN.CENTER)

page_num(s, 9)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 10 — Methodology  (split: 10.jpg auditorium)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 10)
add_logo(s); label(s, "METHODOLOGY")

tx(s, Inches(1.0), Inches(1.1), Inches(6.5), Inches(0.5),
   "The Zero \u2192 One Framework", 28, TEXT, bold=True, font=H)

phases = [
    ("Validate", "Weeks 1\u20133", "Problem clarity\nMarket validation\nUser interviews"),
    ("Design & Build", "Weeks 4\u20138", "Wireframes\nPrototyping\nMVP development"),
    ("Launch & Traction", "Weeks 9\u201312", "User acquisition\nFeedback loops\nIteration & Demo Day"),
]
pw = Inches(2.1); ph = Inches(4.5); pgap = Inches(0.15); sy = Inches(1.8)
for i, (title, weeks, desc) in enumerate(phases):
    x = Inches(1.0) + (pw + pgap) * i

    card(s, x, sy, pw, ph)
    # Amber top bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, sy, pw, Inches(0.04))
    bar.fill.solid(); bar.fill.fore_color.rgb = AMBER; bar.line.fill.background()

    # Phase number circle
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), sy + Inches(0.25), Inches(0.45), Inches(0.45))
    circ.fill.solid(); circ.fill.fore_color.rgb = AMBER; circ.line.fill.background()
    tx(s, x + Inches(0.15), sy + Inches(0.28), Inches(0.45), Inches(0.4),
       str(i+1), 16, WHITE, bold=True, font=H, align=PP_ALIGN.CENTER)

    tx(s, x + Inches(0.7), sy + Inches(0.3), pw - Inches(0.85), Inches(0.3),
       title, 13, TEXT, bold=True, font=H)

    # Weeks badge
    acard(s, x + Inches(0.15), sy + Inches(0.9), Inches(1.1), Inches(0.3))
    tx(s, x + Inches(0.15), sy + Inches(0.92), Inches(1.1), Inches(0.25),
       weeks, 9, DARK_AMBER, bold=True, font=H, align=PP_ALIGN.CENTER)

    tx(s, x + Inches(0.15), sy + Inches(1.4), pw - Inches(0.3), Inches(2.8),
       desc, 12, SUBTEXT, font=B, sp=1.7)

    # Arrow between phases
    if i < 2:
        tx(s, x + pw + Inches(0.0), sy + ph / 2 - Inches(0.1), Inches(0.15), Inches(0.2),
           "\u25B6", 10, AMBER, font=B, align=PP_ALIGN.CENTER)

page_num(s, 10)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 11 — Program Outcomes  (split: 11.jpg rocket)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, WHITE); bar_top(s); right_image(s, 11)
add_logo(s); label(s, "PROGRAM OUTCOMES")

tx(s, Inches(1.0), Inches(1.1), Inches(6.5), Inches(0.6),
   "By the end of 12 weeks,\nfounders will have:", 22, TEXT, bold=True, font=H, sp=1.2)

outcomes = [
    ("A working MVP", "Functional product solving a real problem."),
    ("Validated problem-solution fit", "Evidence your solution matches market need."),
    ("Early users or customers", "Real traction with actual users."),
    ("Pitch-ready narrative", "Compelling story for investors and partners."),
    ("Global pitch platform", "Demo Day exposure to a worldwide audience."),
    ("Next-stage roadmap", "Clear plan for scaling beyond the program."),
]
cw = Inches(3.1); ch = Inches(0.95); gx = Inches(0.2); gy = Inches(0.1)
sx = Inches(1.0); sy = Inches(2.2)
for i, (title, desc) in enumerate(outcomes):
    col = i % 2; row = i // 2
    x = sx + (cw + gx) * col; y = sy + (ch + gy) * row
    card(s, x, y, cw, ch)

    # Checkmark circle
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.12), y + Inches(0.12), Inches(0.3), Inches(0.3))
    circ.fill.solid(); circ.fill.fore_color.rgb = AMBER_50
    circ.line.color.rgb = AMBER_200; circ.line.width = Pt(1)
    tx(s, x + Inches(0.12), y + Inches(0.12), Inches(0.3), Inches(0.3),
       "\u2714", 11, DARK_AMBER, bold=True, font=B, align=PP_ALIGN.CENTER)

    tx(s, x + Inches(0.5), y + Inches(0.1), cw - Inches(0.65), Inches(0.3),
       title, 12, TEXT, bold=True, font=H)
    tx(s, x + Inches(0.15), y + Inches(0.5), cw - Inches(0.3), Inches(0.35),
       desc, 10, SUBTEXT, font=B)

page_num(s, 11)


# ═══════════════════════════════════════════════════════════════
#  SLIDE 12 — Why Z1 Forge Exists  (full-bleed: 12.jpg handshake)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, BLACK)
full_bg(s, 12)
overlay(s, 0.60)
bar_top(s)
add_logo(s, Inches(0.9), Inches(0.5), Inches(0.45))
label(s, "WHY Z1 FORGE EXISTS", top=Inches(0.6), color=AMBER)

tx(s, Inches(0.9), Inches(1.6), Inches(10), Inches(1.0),
   "The startup ecosystem often celebrates\nideas and funding.", 30, WHITE, bold=True, font=H, sp=1.3)

tx(s, Inches(0.9), Inches(3.0), Inches(10), Inches(0.5),
   "We focus on the most critical stage:", 18, LIGHT_TEXT, font=B)

# Highlighted amber-border box
rect = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(0.9), Inches(3.8), Inches(10.5), Inches(1.5))
rect.fill.solid(); rect.fill.fore_color.rgb = RGBColor(0x1A, 0x14, 0x05)
_alpha(rect, 0.80)
rect.line.color.rgb = AMBER; rect.line.width = Pt(2)
try: rect.adjustments[0] = 0.02
except: pass

tx(s, Inches(1.3), Inches(4.0), Inches(9.6), Inches(1.1),
   "The messy, uncertain, high-risk\nZero \u2192 One phase.",
   28, AMBER, bold=True, font=H, align=PP_ALIGN.CENTER, sp=1.3)

tx(s, Inches(0.9), Inches(5.8), Inches(10), Inches(0.5),
   "Z1 Forge exists to increase the probability that founders actually launch.",
   16, WHITE, bold=True, font=H)

page_num(s, 12, RGBColor(0x88, 0x88, 0x88))


# ═══════════════════════════════════════════════════════════════
#  SLIDE 13 — Call to Action  (full-bleed: 13.jpg molten gold)
# ═══════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, BLACK)
full_bg(s, 13)
overlay(s, 0.55)
bar_top(s)
add_logo(s, Inches(0.9), Inches(0.5), Inches(0.45))

tx(s, Inches(0.9), Inches(1.6), Inches(10), Inches(1.0),
   "Build. Ship. Validate. Iterate.", 44, AMBER, bold=True, font=H)

rule(s, Inches(0.9), Inches(3.0), Inches(4), AMBER)

tx(s, Inches(0.9), Inches(3.4), Inches(9), Inches(1.2),
   "If you\u2019re serious about turning your idea\ninto a real product in 12 weeks \u2014\nZ1 Forge is your launchpad.",
   20, WHITE, font=B, sp=1.5)

# CTA button
btn = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                          Inches(0.9), Inches(5.2), Inches(3.5), Inches(0.8))
btn.fill.solid(); btn.fill.fore_color.rgb = AMBER; btn.line.fill.background()
try: btn.adjustments[0] = 0.15
except: pass
tx(s, Inches(0.9), Inches(5.28), Inches(3.5), Inches(0.6),
   "Apply Now  \u2192", 20, BLACK, bold=True, font=H, align=PP_ALIGN.CENTER)

tx(s, Inches(0.9), Inches(6.2), Inches(5), Inches(0.3),
   "z1.sbv.one", 14, LIGHT_TEXT, font=B)

# Bottom amber strip
bb = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.05), W, Inches(0.45))
bb.fill.solid(); bb.fill.fore_color.rgb = AMBER; bb.line.fill.background()
tx(s, Inches(0.9), Inches(7.1), Inches(8), Inches(0.35),
   "Z1 Forge by SBV  \u2022  From Idea to MVP in 12 Weeks", 11, BLACK, bold=True, font=H)
page_num(s, 13, RGBColor(0x66, 0x66, 0x66))


# ═══════════════════════════════════════════════════════════════
out = os.path.join(BASE, "public", "Z1-Forge-Presentation.pptx")
prs.save(out)
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
